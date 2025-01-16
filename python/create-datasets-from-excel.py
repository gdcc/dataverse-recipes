import requests
from easyDataverse import Dataverse
# we use openpyxl because it can pull hyperlinks out of Excel
from openpyxl import load_workbook
import json
import time
import os

#SERVER_URL="https://demo.dataverse.org"
SERVER_URL="http://localhost:8080"
API_TOKEN="3e2b87bf-6652-47f1-b0bb-ddc0aa1a9106"
# The collection where the datasets will be created
COLLECTION="collection1"

try:
    SERVER_URL=os.environ['SERVER_URL']
    print(f"Using SERVER_URL environment variable: {SERVER_URL}")
except:
    print(f"Using SERVER_URL from script: {SERVER_URL}")

try:
    API_TOKEN=os.environ['API_TOKEN']
    print(f"Using API_TOKEN environment variable: {API_TOKEN}")
except:
    print(f"Using API_TOKEN from script: {API_TOKEN}")

try:
    COLLECTION=os.environ['COLLECTION']
    print(f"Using COLLECTION environment variable: {COLLECTION}")
except:
    print(f"Using COLLECTION from script: {COLLECTION}")

wb = load_workbook('data.xlsx')
sheets = wb.sheetnames
ws = wb[sheets[1]]

# Connect to a Dataverse installation
dataverse = Dataverse(
  server_url=SERVER_URL,
  api_token=API_TOKEN,
)

# min_row=2 to skip the header row
count=0
for key, *values in ws.iter_rows(min_row=2):

    # Strangely, datasets created with EasyDataverse have
    # Custom Terms rather than CC0, which is the default,
    # and it's not possible to change the license afterward
    # using EasyDataverse: https://github.com/gdcc/easyDataverse/issues/29
    # So, we create a basic dataset using requests and
    # then update the fields later with EasyDataverse.
    with open('initial-dataset.json') as f:
        r = requests.post(SERVER_URL + "/api/dataverses/" + COLLECTION + "/datasets", data=f, headers={"X-Dataverse-key": API_TOKEN})
        # We'll use this pid later to update the dataset
        pid = r.json()['data']['persistentId']

    # Get the fields from Excel
    title = values[0].value
    title = title.replace("’","'").replace("–","-")
    print("title: " + title)
    description = values[1].value
    print("description: " + description)
    agency_responsible = values[2].value
    print("agency responsible: " + agency_responsible)
    access_link_name = str(values[3].value)
    access_link_name = access_link_name.replace("’","'").replace("–","-")
    print("access link name: " + access_link_name)
    access_link_url = values[3].hyperlink.display
    print("access link url: " + access_link_url)

    # Load the dataset we created earlier
    dataset = dataverse.load_dataset(pid)

    # Update the metadata fields
    dataset.citation.title = title
    dataset.citation.subject = ["Other"]
    dataset.metadatablocks["citation"].ds_description[0].value = description
    dataset.metadatablocks["citation"].author[0].name = agency_responsible
    dataset.metadatablocks["citation"].author[0].affiliation = None
    dataset.metadatablocks["citation"].dataset_contact[0].name = "Harvard Dataverse Support"
    dataset.metadatablocks["citation"].dataset_contact[0].email = "support@dataverse.harvard.edu"
    dataset.metadatablocks["citation"].origin_of_sources = '<a href="' + access_link_url + '">' + access_link_name + '</a>'

    # Update the dataset, pushing the new fields to the server
    dataset.update()

    # Sleep a bit before creating the next dataset
    time.sleep(2)

    # Uncomment this if you want to test first by only creating a single dataset
    # count += 1
    # if (count > 0):
    #    exit()
